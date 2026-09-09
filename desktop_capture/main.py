from __future__ import annotations

import math
import queue
import csv
from pathlib import Path
import re
import shutil
import threading
import time
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import winreg

from adb_camera import AdbCamera, AdbCameraError
from dataset_manifest import append_manifest, successful_shots

try:
    import serial
    from serial.tools import list_ports
except ImportError:
    serial = None
    list_ports = None


DEFAULTS = {
    "HMAX": 100.0, "HC": 25.0, "BR": 0.65, "BS": 1.20,
    "LOOP_HZ": 195.3125, "X_GAIN": 1.0, "Y_GAIN": 1.0,
    "X_OFFSET": 0.0, "Y_OFFSET": 0.0, "XY_COUPLING": 0.0,
    "ASYMMETRY": 0.0, "AMPLITUDE": 1500.0,
}
RANGES = {
    "HMAX": (20, 500), "HC": (1, 100), "BR": (0.05, 1.5),
    "BS": (0.1, 2.0), "LOOP_HZ": (20, 400),
    "X_GAIN": (0.1, 2.0), "Y_GAIN": (0.1, 2.0),
    "X_OFFSET": (-0.8, 0.8), "Y_OFFSET": (-0.8, 0.8),
    "XY_COUPLING": (-0.5, 0.5), "ASYMMETRY": (-0.3, 0.3),
    "AMPLITUDE": (200, 1900),
}
PARAM_INFO = {
    "HMAX": ("最大磁场 Hmax", "A/m"), "HC": ("矫顽力 Hc", "A/m"),
    "BR": ("剩磁 Br", "T"), "BS": ("饱和磁感应 Bs", "T"),
    "LOOP_HZ": ("回线刷新率", "Hz"), "X_GAIN": ("X 增益", "倍"),
    "Y_GAIN": ("Y 增益", "倍"), "X_OFFSET": ("X 偏置", "满量程"),
    "Y_OFFSET": ("Y 偏置", "满量程"), "XY_COUPLING": ("轴间耦合", "比例"),
    "ASYMMETRY": ("上下支不对称", "比例"), "AMPLITUDE": ("DAC 幅度", "LSB"),
}


def build_apply_commands(values: dict[str, float]) -> list[str]:
    """Build a transition that remains valid after every legacy SET command.

    Firmware HYST-V1 validates relational constraints on each SET, rather than
    only on APPLY. Stage wide HMAX/BS bounds first, then set the dependent HC/BR
    before their final bounds so transitions between any two valid rows work.
    """
    ordered_keys = (
        "HC", "HMAX", "BR", "BS", "LOOP_HZ", "X_GAIN", "Y_GAIN",
        "X_OFFSET", "Y_OFFSET", "XY_COUPLING", "ASYMMETRY", "AMPLITUDE",
    )
    lines = ["SET HMAX 500.00000", "SET BS 2.00000"]
    lines.extend(
        f"SET {key} {int(values[key]) if key == 'AMPLITUDE' else f'{values[key]:.5f}'}"
        for key in ordered_keys
    )
    lines.append("APPLY")
    return lines


def build_setall_command(values: dict[str, float]) -> str:
    return (
        "SETALL "
        f"{values['HMAX']:.5f} {values['HC']:.5f} {values['BR']:.5f} {values['BS']:.5f} "
        f"{values['LOOP_HZ']:.5f} {values['X_GAIN']:.5f} {values['Y_GAIN']:.5f} "
        f"{values['X_OFFSET']:.5f} {values['Y_OFFSET']:.5f} "
        f"{values['XY_COUPLING']:.5f} {values['ASYMMETRY']:.5f} {int(values['AMPLITUDE'])}"
    )


class HysteresisUI(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("STM32F407 磁滞回线控制台")
        self.geometry("1260x850")
        self.minsize(1120, 760)
        self.ser = None
        self.rx_queue: queue.Queue[str] = queue.Queue()
        self.tx_queue: queue.Queue[bytes] = queue.Queue()
        self.stop_event = threading.Event()
        self.handshake_ok = False
        self.setall_supported = False
        self.caps_pending = False
        self.caps_received = False
        self.handshake_job = None
        self.port_devices: dict[str, str] = {}
        self.apply_job = None
        self.vars = {k: tk.DoubleVar(value=v) for k, v in DEFAULTS.items()}
        self.status = tk.StringVar(value="未连接")
        self.port_var = tk.StringVar()
        self.auto_apply = tk.BooleanVar(value=True)
        self.table_rows: list[dict[str, float | str]] = []
        self.table_index = -1
        self.table_waiting = False
        self.table_running = False
        self.table_timer = None
        self.status_retry_job = None
        self.apply_watchdog_job = None
        self.table_apply_attempts = 0
        self.capture_recovery_job = None
        self.capture_recovery_attempts = 0
        self.table_interval = tk.DoubleVar(value=3.0)
        self.table_mode = tk.StringVar(value="手动")
        self.table_file = tk.StringVar(value="未加载参数表")
        self.table_position = tk.StringVar(value="0 / 0")
        self.table_current = tk.StringVar(value="当前参数：—")
        bundled_adb = Path(r"D:\Program Files\platform-tools\adb.exe")
        self.adb_path = tk.StringVar(value=shutil.which("adb") or (str(bundled_adb) if bundled_adb.exists() else "adb"))
        self.adb_device = tk.StringVar()
        self.adb_devices: dict[str, str] = {}
        self.capture_output = tk.StringVar(value=str(Path.cwd() / "dataset"))
        self.capture_settle = tk.DoubleVar(value=0.8)
        self.capture_count = tk.IntVar(value=1)
        self.capture_retries = tk.IntVar(value=2)
        self.capture_shutter = tk.StringVar(value="相机键")
        self.capture_tap_x = tk.IntVar(value=0)
        self.capture_tap_y = tk.IntVar(value=0)
        self.capture_status = tk.StringVar(value="ADB 未检测")
        self.capture_running = False
        self.capture_previous_auto_apply = True
        self.capture_stop = threading.Event()
        self.capture_completed: dict[str, set[int]] = {}
        self._build()
        self.refresh_ports()
        self.after(50, self.poll_rx)
        self.after(100, self.draw_preview)
        self.protocol("WM_DELETE_WINDOW", self.close)

    def _build(self) -> None:
        top = ttk.Frame(self, padding=10)
        top.pack(fill="x")
        ttk.Label(top, text="蓝牙串口").pack(side="left")
        self.port_box = ttk.Combobox(top, textvariable=self.port_var, width=22, state="readonly")
        self.port_box.pack(side="left", padx=6)
        ttk.Button(top, text="刷新", command=self.refresh_ports).pack(side="left")
        self.connect_btn = ttk.Button(top, text="连接", command=self.toggle_connection)
        self.connect_btn.pack(side="left", padx=6)
        ttk.Button(top, text="PING", command=lambda: self.send("PING")).pack(side="left")
        ttk.Button(top, text="读取状态", command=lambda: self.send("GET STATUS")).pack(side="left", padx=6)
        ttk.Label(top, textvariable=self.status).pack(side="right")
        ttk.Label(top, textvariable=self.capture_status, foreground="#0b7285").pack(side="right", padx=18)

        body = ttk.Panedwindow(self, orient="horizontal")
        body.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        controls_outer = ttk.Frame(body)
        controls_canvas = tk.Canvas(controls_outer, highlightthickness=0, background="#f0f0f0")
        controls_scroll = ttk.Scrollbar(controls_outer, orient="vertical", command=controls_canvas.yview)
        controls_canvas.configure(yscrollcommand=controls_scroll.set)
        controls_scroll.pack(side="right", fill="y")
        controls_canvas.pack(side="left", fill="both", expand=True)
        controls = ttk.Frame(controls_canvas, padding=8)
        controls_window = controls_canvas.create_window((0, 0), window=controls, anchor="nw")
        controls.bind("<Configure>", lambda _e: controls_canvas.configure(scrollregion=controls_canvas.bbox("all")))
        controls_canvas.bind("<Configure>", lambda e: controls_canvas.itemconfigure(controls_window, width=e.width))
        controls_canvas.bind("<Enter>", lambda _e: controls_canvas.bind_all(
            "<MouseWheel>", lambda event: controls_canvas.yview_scroll(int(-event.delta / 120), "units")))
        controls_canvas.bind("<Leave>", lambda _e: controls_canvas.unbind_all("<MouseWheel>"))
        right = ttk.Frame(body, padding=8)
        body.add(controls_outer, weight=2)
        body.add(right, weight=3)

        for row, key in enumerate(DEFAULTS):
            lo, hi = RANGES[key]
            label, unit = PARAM_INFO[key]
            ttk.Label(controls, text=label, width=17).grid(row=row, column=0, sticky="w", pady=3)
            scale = ttk.Scale(controls, from_=lo, to=hi, variable=self.vars[key], command=lambda _v, k=key: self.changed(k))
            scale.grid(row=row, column=1, sticky="ew", padx=6)
            spin = ttk.Spinbox(controls, from_=lo, to=hi, textvariable=self.vars[key], width=10,
                               increment=1 if key == "AMPLITUDE" else 0.01, command=lambda k=key: self.changed(k))
            spin.grid(row=row, column=2)
            ttk.Label(controls, text=unit, width=7, foreground="#52606d").grid(row=row, column=3, sticky="w", padx=(5, 0))
            spin.bind("<Return>", lambda _e, k=key: self.changed(k))
            spin.bind("<FocusOut>", lambda _e, k=key: self.changed(k))
        controls.columnconfigure(1, weight=1)

        actions = ttk.Frame(controls)
        actions.grid(row=len(DEFAULTS), column=0, columnspan=4, sticky="ew", pady=10)
        for name in ("SOFT", "HARD", "UNSAT"):
            ttk.Button(actions, text=name, command=lambda n=name: self.preset(n)).pack(side="left", padx=3)
        ttk.Button(actions, text="恢复默认", command=self.restore_defaults).pack(side="left", padx=8)
        ttk.Button(actions, text="立即应用", command=self.apply_all).pack(side="right")
        ttk.Checkbutton(controls, text="参数变化后自动应用（120 ms 防抖）", variable=self.auto_apply).grid(
            row=len(DEFAULTS)+1, column=0, columnspan=4, sticky="w")

        table = ttk.LabelFrame(controls, text="参数表模式（只切换参数，不会拍照）", padding=8)
        table.grid(row=len(DEFAULTS)+2, column=0, columnspan=4, sticky="ew", pady=(10, 0))
        ttk.Button(table, text="加载 CSV / Excel", command=self.load_parameter_table).grid(row=0, column=0, columnspan=2, sticky="w")
        ttk.Label(table, textvariable=self.table_file).grid(row=0, column=2, columnspan=3, sticky="w", padx=8)
        ttk.Radiobutton(table, text="手动", variable=self.table_mode, value="手动", command=self.table_mode_changed).grid(row=1, column=0)
        ttk.Radiobutton(table, text="定时", variable=self.table_mode, value="定时", command=self.table_mode_changed).grid(row=1, column=1)
        ttk.Label(table, text="间隔(s)").grid(row=1, column=2, padx=(10, 2))
        ttk.Spinbox(table, from_=0.5, to=3600, increment=0.5, width=7, textvariable=self.table_interval).grid(row=1, column=3)
        self.table_start_btn = ttk.Button(table, text="开始定时（不拍照）", command=self.toggle_table_timer)
        self.table_start_btn.grid(row=1, column=4, padx=5)
        ttk.Button(table, text="上一条", command=lambda: self.step_table(-1)).grid(row=2, column=0, pady=6)
        ttk.Button(table, text="下一条", command=lambda: self.step_table(1)).grid(row=2, column=1, pady=6)
        ttk.Label(table, textvariable=self.table_position).grid(row=2, column=2, columnspan=2)
        ttk.Label(table, textvariable=self.table_current, wraplength=520, justify="left",
                  background="#eef6f8", padding=7).grid(row=3, column=0, columnspan=5, sticky="ew", pady=(2, 0))
        table.columnconfigure(4, weight=1)

        capture = ttk.LabelFrame(controls, text="有线 ADB 自动拍照", padding=8)
        capture.grid(row=len(DEFAULTS)+3, column=0, columnspan=4, sticky="ew", pady=(10, 0))
        ttk.Label(capture, text="ADB").grid(row=0, column=0, sticky="w")
        ttk.Entry(capture, textvariable=self.adb_path, width=38).grid(row=0, column=1, columnspan=2, sticky="ew", padx=5)
        ttk.Button(capture, text="浏览", command=self.choose_adb).grid(row=0, column=3)
        ttk.Button(capture, text="检测手机", command=self.refresh_adb_devices).grid(row=0, column=4, padx=(5, 0))
        ttk.Label(capture, text="手机").grid(row=1, column=0, sticky="w", pady=5)
        self.adb_device_box = ttk.Combobox(capture, textvariable=self.adb_device, state="readonly", width=36)
        self.adb_device_box.grid(row=1, column=1, columnspan=2, sticky="ew", padx=5)
        ttk.Button(capture, text="打开相机", command=self.open_phone_camera).grid(row=1, column=3)
        ttk.Button(capture, text="测试拍照", command=self.test_phone_capture).grid(row=1, column=4, padx=(5, 0))
        ttk.Label(capture, text="输出").grid(row=2, column=0, sticky="w")
        ttk.Entry(capture, textvariable=self.capture_output).grid(row=2, column=1, columnspan=2, sticky="ew", padx=5)
        ttk.Button(capture, text="浏览", command=self.choose_capture_output).grid(row=2, column=3)
        self.capture_start_btn = ttk.Button(capture, text="▶ 开始自动拍照", command=self.start_dataset_capture)
        self.capture_start_btn.grid(row=2, column=4, padx=(5, 0))
        ttk.Button(capture, text="停止", command=self.stop_dataset_capture).grid(row=2, column=5, padx=(5, 0))
        ttk.Label(capture, text="稳定(s)").grid(row=3, column=0, sticky="w", pady=5)
        ttk.Spinbox(capture, from_=0.1, to=30, increment=0.1, width=6, textvariable=self.capture_settle).grid(row=3, column=1, sticky="w", padx=5)
        ttk.Label(capture, text="每组张数").grid(row=3, column=2, sticky="e")
        ttk.Spinbox(capture, from_=1, to=20, width=5, textvariable=self.capture_count).grid(row=3, column=3, sticky="w", padx=5)
        ttk.Label(capture, text="重试").grid(row=3, column=4, sticky="e")
        ttk.Spinbox(capture, from_=0, to=10, width=5, textvariable=self.capture_retries).grid(row=3, column=5, sticky="w")
        ttk.Label(capture, text="快门").grid(row=4, column=0, sticky="w")
        ttk.Combobox(capture, textvariable=self.capture_shutter, values=("相机键", "屏幕坐标"),
                     state="readonly", width=8).grid(row=4, column=1, sticky="w", padx=5)
        ttk.Label(capture, text="X").grid(row=4, column=2, sticky="e")
        ttk.Spinbox(capture, from_=0, to=10000, width=6, textvariable=self.capture_tap_x).grid(row=4, column=3, sticky="w", padx=5)
        ttk.Label(capture, text="Y").grid(row=4, column=4, sticky="e")
        ttk.Spinbox(capture, from_=0, to=10000, width=6, textvariable=self.capture_tap_y).grid(row=4, column=5, sticky="w")
        ttk.Label(capture, textvariable=self.capture_status, foreground="#0b7285").grid(
            row=5, column=0, columnspan=6, sticky="w", pady=(4, 0))
        capture.columnconfigure(2, weight=1)

        self.canvas = tk.Canvas(right, bg="#101820", highlightthickness=0)
        self.canvas.pack(fill="both", expand=True)
        ttk.Label(right, text="设备日志").pack(anchor="w", pady=(8, 2))
        log_frame = ttk.Frame(right)
        log_frame.pack(fill="x")
        self.log = tk.Text(log_frame, height=9, state="disabled", font=("Consolas", 10), wrap="none")
        log_scroll_y = ttk.Scrollbar(log_frame, orient="vertical", command=self.log.yview)
        log_scroll_x = ttk.Scrollbar(log_frame, orient="horizontal", command=self.log.xview)
        self.log.configure(yscrollcommand=log_scroll_y.set, xscrollcommand=log_scroll_x.set)
        self.log.grid(row=0, column=0, sticky="nsew")
        log_scroll_y.grid(row=0, column=1, sticky="ns")
        log_scroll_x.grid(row=1, column=0, sticky="ew")
        log_frame.columnconfigure(0, weight=1)
        self.canvas.bind("<Configure>", lambda _e: self.draw_preview())

    def choose_adb(self) -> None:
        path = filedialog.askopenfilename(title="选择 adb.exe", filetypes=[("adb", "adb.exe"), ("程序", "*.exe")])
        if path:
            self.adb_path.set(path)

    def choose_capture_output(self) -> None:
        path = filedialog.askdirectory(title="选择数据集输出目录", initialdir=self.capture_output.get())
        if path:
            self.capture_output.set(path)

    def refresh_adb_devices(self) -> None:
        self.capture_status.set("正在检测 ADB 手机…")
        adb = self.adb_path.get().strip() or "adb"

        def worker() -> None:
            try:
                devices = AdbCamera.list_devices(adb, timeout=8.0)
                self.after(0, lambda: self._show_adb_devices(devices))
            except Exception as exc:
                self.after(0, lambda: self.capture_status.set(f"ADB 检测失败：{exc}"))

        threading.Thread(target=worker, daemon=True).start()

    def _show_adb_devices(self, devices: list[tuple[str, str]]) -> None:
        self.adb_devices.clear()
        labels = []
        for serial_number, state in devices:
            label = f"{serial_number} — {state}"
            labels.append(label)
            self.adb_devices[label] = serial_number
        self.adb_device_box["values"] = labels
        ready = [label for label in labels if label.endswith("— device")]
        if ready:
            self.adb_device.set(ready[0])
            self.capture_status.set(f"已连接：{self.adb_devices[ready[0]]}")
        elif labels:
            self.adb_device.set(labels[0])
            state = labels[0].rsplit("—", 1)[-1].strip()
            self.capture_status.set("手机未就绪：" + ("请在手机上允许 USB 调试" if state == "unauthorized" else state))
        else:
            self.adb_device.set("")
            self.capture_status.set("未发现手机；请连接 USB 数据线并开启 USB 调试")

    def _phone(self, cancellable: bool = False) -> AdbCamera:
        label = self.adb_device.get()
        serial_number = self.adb_devices.get(label, label.split(" — ", 1)[0].strip())
        if not serial_number:
            raise AdbCameraError("请先点击“检测手机”并选择设备")
        return AdbCamera(self.adb_path.get().strip() or "adb", serial_number, timeout=15.0,
                         cancel_event=self.capture_stop if cancellable else None)

    def _run_phone_action(self, status: str, action, done=None) -> None:
        self.capture_status.set(status)

        def worker() -> None:
            try:
                result = action()
                self.after(0, lambda: done(result) if done else self.capture_status.set("操作成功"))
            except Exception as exc:
                self.after(0, lambda: self.capture_status.set(f"操作失败：{exc}"))

        threading.Thread(target=worker, daemon=True).start()

    def open_phone_camera(self) -> None:
        try:
            phone = self._phone()
        except Exception as exc:
            self.capture_status.set(str(exc)); return
        self._run_phone_action("正在打开手机相机…", phone.open_camera,
                               lambda _r: self.capture_status.set("相机已打开，请确认取景和曝光"))

    def test_phone_capture(self) -> None:
        try:
            phone = self._phone()
            output = Path(self.capture_output.get()).expanduser().resolve() / "test"
        except Exception as exc:
            self.capture_status.set(str(exc)); return

        def action():
            output.mkdir(parents=True, exist_ok=True)
            phone.open_camera()
            time.sleep(1.0)
            return phone.capture(output / f"adb_test_{time.strftime('%Y%m%d_%H%M%S')}.jpg", tap=tap)

        tap = self._shutter_tap()
        self._run_phone_action("正在测试拍照并回传…", action,
                               lambda result: self.capture_status.set(f"测试成功：{result[0]}"))

    def _shutter_tap(self) -> tuple[int, int] | None:
        if self.capture_shutter.get() != "屏幕坐标":
            return None
        return max(0, int(self.capture_tap_x.get())), max(0, int(self.capture_tap_y.get()))

    @staticmethod
    def _sample_id(item: dict[str, float | str], index: int) -> str:
        raw = str(item.get("ID") or f"HYS_{index + 1:06d}").strip()
        safe = re.sub(r"[^0-9A-Za-z_.-]+", "_", raw).strip("._")
        return safe or f"HYS_{index + 1:06d}"

    def start_dataset_capture(self) -> None:
        if self.capture_running:
            return
        if not self.table_rows:
            messagebox.showinfo("自动拍照", "请先加载 CSV 或 Excel 参数表")
            return
        if not self.ser or not self.handshake_ok:
            messagebox.showinfo("自动拍照", "请先连接 STM32，并确认日志已收到 PONG")
            return
        if not self.setall_supported:
            messagebox.showerror(
                "固件版本过旧，禁止自动采集",
                "当前 STM32 没有启用 SETALL 原子协议。旧协议每组发送约 15 条命令，"
                "会随机丢失单个参数，不能用于无人值守采集。\n\n"
                "请烧录 Output/atk_f407.hex，重新连接后确认日志出现：\n"
                "CAPS SETALL HYST-V2\n"
                "已启用 SETALL 原子协议：每组参数仅发送 1 条命令",
            )
            self.capture_status.set("禁止启动：请先烧录 HYST-V2 SETALL 固件")
            return
        try:
            phone = self._phone()
            root = Path(self.capture_output.get()).expanduser().resolve()
            root.mkdir(parents=True, exist_ok=True)
            (root / "images").mkdir(exist_ok=True)
            count = max(1, int(self.capture_count.get()))
            if float(self.capture_settle.get()) < 0:
                raise ValueError("稳定等待不能小于 0")
            sample_ids = [self._sample_id(item, i) for i, item in enumerate(self.table_rows)]
            if len(sample_ids) != len(set(sample_ids)):
                raise ValueError("参数表中的 SAMPLE_ID 经文件名规范化后存在重复，请先改为唯一编号")
            self._write_parameters_snapshot(root)
        except Exception as exc:
            messagebox.showerror("无法开始", str(exc)); return

        self.stop_table_timer()
        self.capture_running = True
        self.capture_previous_auto_apply = self.auto_apply.get()
        self.auto_apply.set(False)
        # A fresh event per run prevents a quickly restarted run from clearing
        # the cancellation signal still owned by the previous worker.
        self.capture_stop = threading.Event()
        run_stop = self.capture_stop
        self.capture_start_btn.config(state="disabled", text="自动拍照运行中…")
        self.capture_completed = successful_shots(root / "manifest.csv", root)
        target = next((i for i, item in enumerate(self.table_rows)
                       if len(self.capture_completed.get(self._sample_id(item, i), set())) < count), None)
        if target is None:
            self._finish_dataset_capture("参数表中的照片均已采集，无需续拍")
            return
        done_count = sum(1 for i, item in enumerate(self.table_rows)
                         if len(self.capture_completed.get(self._sample_id(item, i), set())) >= count)
        self.capture_status.set(f"正在准备手机；将从第 {target + 1} 行续拍（已完成 {done_count} 组）")

        def prepare() -> None:
            try:
                phone.open_camera()
                self.after(0, lambda: self.apply_table_index(target)
                           if self.capture_running and self.capture_stop is run_stop else None)
            except Exception as exc:
                self.after(0, lambda: self._finish_dataset_capture(f"手机准备失败：{exc}"))

        threading.Thread(target=prepare, daemon=True).start()

    def _write_parameters_snapshot(self, root: Path) -> None:
        columns = ["sample_id"]
        for item in self.table_rows:
            for key in item:
                if key not in ("_ROW", "ID") and key not in columns:
                    columns.append(key)
        with (root / "parameters.csv").open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=columns)
            writer.writeheader()
            for index, item in enumerate(self.table_rows):
                writer.writerow({"sample_id": self._sample_id(item, index),
                                 **{key: value for key, value in item.items() if key not in ("_ROW", "ID")}})

    def stop_dataset_capture(self) -> None:
        if self.capture_running:
            self.capture_stop.set()
            self.capture_running = False
            self.table_waiting = False
            self.stop_table_timer()
            if self.status_retry_job:
                self.after_cancel(self.status_retry_job)
                self.status_retry_job = None
            if self.apply_watchdog_job:
                self.after_cancel(self.apply_watchdog_job)
                self.apply_watchdog_job = None
            if self.capture_recovery_job:
                self.after_cancel(self.capture_recovery_job)
                self.capture_recovery_job = None
            self.auto_apply.set(self.capture_previous_auto_apply)
            self.capture_start_btn.config(state="normal", text="▶ 开始自动拍照")
            self.capture_status.set("自动采集已停止，可稍后断点续拍")
            self.append_log("自动采集已立即停止")

    def _capture_current_row(self) -> None:
        run_stop = self.capture_stop
        if not self.capture_running or run_stop.is_set():
            self._finish_dataset_capture("自动采集已停止")
            return
        index = self.table_index
        item = self.table_rows[index]
        sample_id = self._sample_id(item, index)
        count = max(1, int(self.capture_count.get()))
        completed = self.capture_completed.get(sample_id, set())
        missing = [shot for shot in range(1, count + 1) if shot not in completed]
        root = Path(self.capture_output.get()).expanduser().resolve()
        retries = max(0, int(self.capture_retries.get()))
        tap = self._shutter_tap()
        try:
            phone = self._phone(cancellable=True)
        except Exception as exc:
            self._finish_dataset_capture(f"手机不可用：{exc}")
            return
        self.capture_status.set(f"拍摄 {index + 1}/{len(self.table_rows)}：{sample_id}")

        def worker() -> None:
            error = ""
            for shot in missing:
                if run_stop.is_set():
                    break
                stem = sample_id if count == 1 else f"{sample_id}_{shot:02d}"
                for attempt in range(retries + 1):
                    try:
                        destination, media = phone.capture(root / "images" / f"{stem}.jpg", tap=tap)
                        relative = destination.relative_to(root).as_posix()
                        row = {**item, "sample_id": sample_id, "shot": shot, "status": "ok",
                               "image": relative, "device_serial": phone.serial,
                               "phone_name": media.display_name, "error": ""}
                        append_manifest(root / "manifest.csv", row)
                        self.capture_completed.setdefault(sample_id, set()).add(shot)
                        error = ""
                        break
                    except Exception as exc:
                        error = str(exc)
                        if run_stop.is_set():
                            error = ""
                            break
                        if attempt < retries:
                            if run_stop.wait(0.7):
                                error = ""
                                break
                if error:
                    append_manifest(root / "manifest.csv", {**item, "sample_id": sample_id, "shot": shot,
                                    "status": "error", "device_serial": phone.serial, "error": error})
                    break
            self.after(0, lambda: self._capture_row_done(index, error, run_stop))

        threading.Thread(target=worker, daemon=True).start()

    def _capture_row_done(self, index: int, error: str, run_stop: threading.Event) -> None:
        if run_stop is not self.capture_stop or not self.capture_running:
            return
        if run_stop.is_set():
            self._finish_dataset_capture("自动采集已停止，可稍后断点续拍")
        elif error:
            self.capture_recovery_attempts += 1
            delay = min(15.0, 1.5 * (2 ** min(self.capture_recovery_attempts - 1, 3)))
            self.capture_status.set(
                f"第 {index + 1} 行手机/ADB 暂时失败，{delay:g} 秒后自动恢复：{error}")
            self.append_log(
                f"第 {index + 1} 行拍摄失败，将自动重连并重拍当前组 "
                f"(恢复 {self.capture_recovery_attempts})：{error}")
            self.capture_recovery_job = self.after(
                int(delay * 1000), lambda: self._recover_capture_row(index, run_stop))
        elif index >= len(self.table_rows) - 1:
            self._finish_dataset_capture("自动采集完成")
        else:
            self.capture_recovery_attempts = 0
            self.apply_table_index(index + 1)

    def _recover_capture_row(self, index: int, run_stop: threading.Event) -> None:
        self.capture_recovery_job = None
        if (run_stop is not self.capture_stop or run_stop.is_set()
                or not self.capture_running or self.table_index != index):
            return
        self.capture_status.set(f"正在重连手机并重拍第 {index + 1} 行…")

        def worker() -> None:
            try:
                phone = self._phone(cancellable=True)
                phone.open_camera()
                self.after(0, lambda: self._capture_current_row()
                           if self.capture_running and self.capture_stop is run_stop else None)
            except Exception as exc:
                self.after(0, lambda: self._capture_row_done(index, str(exc), run_stop))

        threading.Thread(target=worker, daemon=True).start()

    def _finish_dataset_capture(self, message: str) -> None:
        self.capture_running = False
        self.capture_stop.set()
        if self.capture_recovery_job:
            self.after_cancel(self.capture_recovery_job)
            self.capture_recovery_job = None
        if self.status_retry_job:
            self.after_cancel(self.status_retry_job)
            self.status_retry_job = None
        if self.apply_watchdog_job:
            self.after_cancel(self.apply_watchdog_job)
            self.apply_watchdog_job = None
        self.auto_apply.set(self.capture_previous_auto_apply)
        self.capture_start_btn.config(state="normal")
        self.capture_start_btn.config(text="▶ 开始自动拍照")
        self.capture_status.set(message)
        self.append_log(message)

    def refresh_ports(self) -> None:
        bluetooth_names = self.windows_bluetooth_port_names()
        entries = []
        self.port_devices.clear()
        if list_ports:
            ports = list(list_ports.comports())
            ports.sort(key=lambda p: ("JDY" not in bluetooth_names.get(p.device, "").upper(), p.device))
            for port in ports:
                hwid = port.hwid.upper()
                known_name = bluetooth_names.get(port.device, "")
                if known_name:
                    direction = known_name
                elif "BTHENUM" in hwid:
                    direction = "蓝牙串口（方向请以 Windows COM 端口页为准）"
                else:
                    direction = "有线串口"
                label = f"{port.device} — {direction} — {port.description}"
                entries.append(label)
                self.port_devices[label] = port.device
        self.port_box["values"] = entries
        if entries and self.port_var.get() not in entries:
            self.port_var.set(entries[0])

    @staticmethod
    def windows_bluetooth_port_names() -> dict[str, str]:
        """Best-effort local aliases; Windows' COM Ports page remains authoritative."""
        # pyserial does not expose the incoming/outgoing direction. Keep known
        # aliases explicit instead of guessing from LOCALMFG, which is not a
        # reliable direction marker on all Windows Bluetooth stacks.
        aliases: dict[str, str] = {}
        try:
            with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, r"HARDWARE\DEVICEMAP\SERIALCOMM") as key:
                index = 0
                while True:
                    try:
                        device, port, _ = winreg.EnumValue(key, index)
                    except OSError:
                        break
                    if str(port).upper() == "COM12": aliases["COM12"] = "传出 — JDY-34S-SPP"
                    if str(port).upper() == "COM13": aliases["COM13"] = "传入 — JDY-34S-SPP"
                    index += 1
        except OSError:
            pass
        return aliases

    def toggle_connection(self) -> None:
        if self.ser:
            self.disconnect()
            return
        if serial is None:
            messagebox.showerror("缺少依赖", "请先运行：python -m pip install -r requirements.txt")
            return
        try:
            label = self.port_var.get()
            port = self.port_devices.get(label, label.split()[0])
            if "传入 — JDY" in label:
                messagebox.showwarning("端口方向错误", "这是 JDY-34 的传入端口。请选择 Windows COM 端口页标记为“传出”的端口。")
                return
            if "蓝牙串口" in label and "JDY" not in label.upper():
                if not messagebox.askyesno(
                    "无法确认设备",
                    "Windows 只把它标成标准串行端口，界面无法确认它属于 JDY-34。\n\n"
                    "如果选到了耳机或其他蓝牙设备，会出现错误 121。仍要尝试吗？",
                ):
                    return
            self.ser = serial.Serial(port, 9600, timeout=0.1, write_timeout=0.25)
            self.stop_event.clear()
            self.handshake_ok = False
            self.setall_supported = False
            self.caps_received = False
            threading.Thread(target=self.reader, daemon=True).start()
            threading.Thread(target=self.writer, daemon=True).start()
            self.status.set(f"正在握手 {port} @ 9600…")
            self.connect_btn.config(text="断开")
            self.send("PING")
            self.caps_pending = True
            self.send("CAPS")
            self.handshake_job = self.after(2500, self.handshake_timeout)
        except Exception as exc:
            self.ser = None
            messagebox.showerror("连接失败", str(exc))

    def disconnect(self) -> None:
        self.stop_dataset_capture()
        self.stop_event.set()
        if self.ser:
            try: self.ser.close()
            except Exception: pass
        self.ser = None
        self.handshake_ok = False
        self.setall_supported = False
        self.caps_received = False
        self.status.set("未连接")
        self.connect_btn.config(text="连接")

    def reader(self) -> None:
        pending = bytearray()
        last_byte_at = 0.0
        while self.ser and not self.stop_event.is_set():
            try:
                waiting = self.ser.in_waiting
                chunk = self.ser.read(waiting if waiting > 0 else 1)
                if not chunk:
                    # HYST-V1 uses a 100 ms HAL transmit timeout. A full STATUS
                    # needs about 177 ms at 9600 baud, so old firmware may stop
                    # mid-line without sending CR/LF. Flush it after a quiet
                    # interval and mark it explicitly as truncated.
                    if pending and time.monotonic() - last_byte_at >= 0.25:
                        line = pending.decode("ascii", errors="replace").strip()
                        pending.clear()
                        if line:
                            self.rx_queue.put("SERIAL PARTIAL: " + line)
                    continue
                pending.extend(chunk)
                last_byte_at = time.monotonic()
                while b"\n" in pending:
                    raw, _, remainder = pending.partition(b"\n")
                    pending = bytearray(remainder)
                    line = raw.rstrip(b"\r").decode("ascii", errors="replace").strip()
                    if line:
                        self.rx_queue.put(line)
                if len(pending) > 4096:
                    self.rx_queue.put("SERIAL ERROR: 接收行超过 4096 字节且没有换行")
                    pending.clear()
            except Exception as exc:
                self.rx_queue.put(f"SERIAL ERROR: {exc}")
                break

    def writer(self) -> None:
        while self.ser and not self.stop_event.is_set():
            try:
                payload = self.tx_queue.get(timeout=0.1)
            except queue.Empty:
                continue
            try:
                self.ser.write(payload)
                self.ser.flush()
                # At 9600 baud the MCU needs a short gap to parse the line
                # and transmit its OK/ERR response without RX overrun.
                time.sleep(0.06)
            except Exception as exc:
                self.rx_queue.put(f"WRITE ERROR: {exc}")
                self.stop_event.set()
                break

    def send(self, line: str) -> None:
        if not self.ser:
            self.append_log("未连接，命令未发送")
            return
        self.send_batch([line])

    def send_batch(self, lines: list[str]) -> None:
        if not self.ser:
            self.append_log("未连接，命令未发送")
            return
        for line in lines:
            self.tx_queue.put((line.strip() + "\n").encode("ascii"))
        self.append_log("> " + (lines[0] if len(lines) == 1 else f"批量发送 {len(lines)} 条命令"))

    def discard_pending_tx(self) -> None:
        discarded = 0
        while True:
            try:
                self.tx_queue.get_nowait()
                discarded += 1
            except queue.Empty:
                break
        if discarded:
            self.append_log(f"已丢弃 {discarded} 条旧的待发送命令，准备重新下发完整事务")

    def handshake_timeout(self) -> None:
        if self.ser and not self.handshake_ok:
            self.status.set("握手失败：端口已打开，但未收到 PONG")
            self.append_log("未收到 PONG。请改选 JDY-34 的传出 COM 口，并确认模块已上电、已烧录最新固件。")

    def changed(self, _key: str) -> None:
        self.draw_preview()
        if self.auto_apply.get():
            if self.apply_job: self.after_cancel(self.apply_job)
            self.apply_job = self.after(120, self.apply_all)

    def apply_all(self) -> None:
        values = {k: v.get() for k, v in self.vars.items()}
        if not (0 < values["HC"] < values["HMAX"] and 0 < values["BR"] < values["BS"]):
            self.append_log("参数错误：需满足 0<HC<HMAX 且 0<BR<BS")
            return
        self.send_parameters(values)
        if self.table_waiting:
            self._arm_apply_watchdog(self.table_index)

    def send_parameters(self, values: dict[str, float]) -> None:
        if self.setall_supported:
            self.send_batch([build_setall_command(values)])
        else:
            self.send_batch(build_apply_commands(values))

    def apply_expected_table_row(self) -> None:
        """Resend the immutable table target, never values echoed by the device."""
        if not self.table_waiting or not (0 <= self.table_index < len(self.table_rows)):
            return
        item = self.table_rows[self.table_index]
        values = dict(DEFAULTS)
        values.update({key: float(value) for key, value in item.items() if key in DEFAULTS})
        self.send_parameters(values)
        self._arm_apply_watchdog(self.table_index)

    def _arm_apply_watchdog(self, expected_index: int) -> None:
        if self.apply_watchdog_job:
            self.after_cancel(self.apply_watchdog_job)
        # 15 commands at 60 ms spacing plus a truncated/full STATUS normally
        # complete in about 1.2 s. Keep margin for Bluetooth scheduling.
        self.apply_watchdog_job = self.after(
            2500, lambda: self._apply_watchdog_expired(expected_index))

    def _apply_watchdog_expired(self, expected_index: int) -> None:
        self.apply_watchdog_job = None
        if not self.table_waiting or not self.ser or self.table_index != expected_index:
            return
        self.append_log(
            f"第 {expected_index + 1} 行在 2.5 秒内未收到可确认的 STATUS，自动恢复事务")
        self.discard_pending_tx()
        self._retry_apply_if_waiting(expected_index)

    def _recover_from_device_error(self, line: str) -> None:
        if not self.table_waiting or not self.ser:
            return
        index = self.table_index
        if self.apply_watchdog_job:
            self.after_cancel(self.apply_watchdog_job)
            self.apply_watchdog_job = None
        if self.status_retry_job:
            self.after_cancel(self.status_retry_job)
        self.discard_pending_tx()
        self.capture_status.set(f"第 {index + 1} 行收到 {line}，正在自动恢复")
        self.append_log(f"设备拒绝了当前事务（{line}），将自动重发第 {index + 1} 行")
        self.status_retry_job = self.after(
            350, lambda expected=index: self._retry_apply_if_waiting(expected))

    def load_parameter_table(self) -> None:
        path = filedialog.askopenfilename(
            title="选择参数表",
            filetypes=[("参数表", "*.csv *.xlsx"), ("CSV", "*.csv"), ("Excel", "*.xlsx")],
        )
        if not path:
            return
        try:
            rows = self.read_parameter_table(Path(path))
            self.validate_parameter_rows(rows)
        except Exception as exc:
            messagebox.showerror("参数表无效", str(exc))
            return
        self.stop_table_timer()
        self.stop_dataset_capture()
        self.table_rows = rows
        self.table_index = -1
        self.table_file.set(f"{Path(path).name}（{len(rows)} 行）")
        self.table_position.set(f"0 / {len(rows)}")
        self.table_current.set("当前参数：尚未输出")
        self.append_log(f"已加载参数表：{path}，共 {len(rows)} 行")

    @staticmethod
    def read_parameter_table(path: Path) -> list[dict[str, float | str]]:
        if path.suffix.lower() == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                records = list(csv.DictReader(handle))
        elif path.suffix.lower() == ".xlsx":
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise RuntimeError("读取 Excel 需要 openpyxl，请重新执行 pip install -r requirements.txt") from exc
            book = load_workbook(path, read_only=True, data_only=True)
            sheet = book.active
            values = list(sheet.iter_rows(values_only=True))
            if not values:
                return []
            headers = [str(v).strip() if v is not None else "" for v in values[0]]
            records = [dict(zip(headers, row)) for row in values[1:] if any(v is not None and str(v).strip() for v in row)]
            book.close()
        else:
            raise ValueError("仅支持 .csv 和 .xlsx")

        normalized = []
        for row_number, record in enumerate(records, start=2):
            item: dict[str, float | str] = {"_ROW": float(row_number)}
            for key, value in record.items():
                name = str(key).strip().upper() if key is not None else ""
                if name in DEFAULTS and value is not None and str(value).strip() != "":
                    item[name] = float(value)
                elif name in ("ID", "NAME", "LABEL", "SAMPLE_ID"):
                    item["ID"] = str(value or "").strip()
                elif name and not name.startswith("_"):
                    item[name] = "" if value is None else str(value).strip()
            normalized.append(item)
        return normalized

    @staticmethod
    def validate_parameter_rows(rows: list[dict[str, float | str]]) -> None:
        if not rows:
            raise ValueError("参数表没有数据行")
        required = ("HMAX", "HC", "BR", "BS")
        for item in rows:
            row = int(float(item["_ROW"]))
            missing = [key for key in required if key not in item]
            if missing:
                raise ValueError(f"第 {row} 行缺少必填参数：{', '.join(missing)}")
            merged = dict(DEFAULTS)
            merged.update({k: float(v) for k, v in item.items() if k in DEFAULTS})
            if not (0 < merged["HC"] < merged["HMAX"] and 0 < merged["BR"] < merged["BS"]):
                raise ValueError(f"第 {row} 行需满足 0<HC<HMAX 且 0<BR<BS")
            for key, (lo, hi) in RANGES.items():
                if not lo <= merged[key] <= hi:
                    raise ValueError(f"第 {row} 行 {key}={merged[key]} 超出范围 [{lo}, {hi}]")

    def apply_table_index(self, index: int) -> None:
        if not self.table_rows or self.table_waiting:
            return
        index = max(0, min(index, len(self.table_rows) - 1))
        item = self.table_rows[index]
        for key, default in DEFAULTS.items():
            self.vars[key].set(float(item.get(key, default)))
        self.table_index = index
        self.table_waiting = True
        self.table_apply_attempts = 1
        self.table_position.set(f"正在应用 {index + 1} / {len(self.table_rows)}")
        self.table_current.set(self.format_table_current(item))
        self.draw_preview()
        self.apply_all()

    def step_table(self, delta: int) -> None:
        if self.capture_running:
            self.append_log("自动拍照运行中，已忽略手动切换")
            return
        if not self.table_rows:
            messagebox.showinfo("参数表", "请先加载 CSV 或 Excel 参数表")
            return
        target = 0 if self.table_index < 0 else self.table_index + delta
        self.apply_table_index(target)

    def table_mode_changed(self) -> None:
        if self.table_mode.get() != "定时":
            self.stop_table_timer()

    def toggle_table_timer(self) -> None:
        if self.capture_running:
            self.append_log("自动拍照运行中，不能启动定时输出")
            return
        if self.table_running:
            self.stop_table_timer()
            return
        if not self.table_rows:
            messagebox.showinfo("参数表", "请先加载参数表")
            return
        self.table_mode.set("定时")
        self.table_running = True
        self.table_start_btn.config(text="停止定时（不拍照）")
        target = 0 if self.table_index < 0 or self.table_index >= len(self.table_rows)-1 else self.table_index + 1
        self.apply_table_index(target)

    def stop_table_timer(self) -> None:
        self.table_running = False
        self.table_start_btn.config(text="开始定时（不拍照）")
        if self.table_timer:
            self.after_cancel(self.table_timer)
            self.table_timer = None

    def schedule_next_table_row(self) -> None:
        if not self.table_running:
            return
        if self.table_index >= len(self.table_rows) - 1:
            self.stop_table_timer()
            self.append_log("参数表定时输出已完成")
            return
        delay = max(0.5, float(self.table_interval.get()))
        self.table_timer = self.after(int(delay * 1000), lambda: self.apply_table_index(self.table_index + 1))

    def format_table_current(self, item: dict[str, float | str]) -> str:
        label = str(item.get("ID") or "未命名")
        values = dict(DEFAULTS)
        values.update({k: float(v) for k, v in item.items() if k in DEFAULTS})
        return (
            f"采集编号：{label}\n"
            f"材料参数：Hmax {values['HMAX']:g} A/m    Hc {values['HC']:g} A/m    "
            f"Br {values['BR']:g} T    Bs {values['BS']:g} T\n"
            f"扫描参数：刷新率 {values['LOOP_HZ']:g} Hz    DAC 幅度 {values['AMPLITUDE']:g} LSB\n"
            f"误差参数：X/Y 增益 {values['X_GAIN']:g}/{values['Y_GAIN']:g} 倍    "
            f"X/Y 偏置 {values['X_OFFSET']:g}/{values['Y_OFFSET']:g} 满量程    "
            f"耦合 {values['XY_COUPLING']:g}    不对称 {values['ASYMMETRY']:g}"
        )

    def preset(self, name: str) -> None:
        # Firmware applies the preset and then returns the new STATUS itself.
        # Sending GET STATUS immediately used to race with APPLY and overwrite
        # the UI with stale values, making one click appear ineffective.
        self.send(f"PRESET {name}")

    def restore_defaults(self) -> None:
        for key, value in DEFAULTS.items(): self.vars[key].set(value)
        self.draw_preview()
        self.apply_all()

    def poll_rx(self) -> None:
        while True:
            try: line = self.rx_queue.get_nowait()
            except queue.Empty: break
            self.append_log("< " + line)
            if line.startswith("SERIAL PARTIAL: STATUS "):
                self.consume_status(line[len("SERIAL PARTIAL: "):], partial=True)
                continue
            if line == "PONG":
                self.handshake_ok = True
                if self.handshake_job:
                    self.after_cancel(self.handshake_job)
                    self.handshake_job = None
                port = self.ser.port if self.ser else ""
                self.status.set(f"设备已连接 {port} @ 9600")
            if line == "CAPS SETALL HYST-V2":
                self.setall_supported = True
                self.caps_received = True
                self.caps_pending = False
                self.append_log("已启用 SETALL 原子协议：每组参数仅发送 1 条命令")
            elif line == "ERR COMMAND" and self.caps_pending and not self.table_waiting:
                self.caps_received = True
                self.caps_pending = False
                self.append_log("当前为旧固件，将使用兼容协议；建议烧录新版 HEX 以提升速度和可靠性")
                continue
            if line.startswith("STATUS "): self.consume_status(line)
            if line.startswith("ERR "):
                self._recover_from_device_error(line)
            if line.startswith("WRITE ERROR") or line.startswith("SERIAL ERROR"):
                self.status.set("串口通信失败")
                self.stop_dataset_capture()
        self.after(50, self.poll_rx)

    def consume_status(self, line: str, partial: bool = False) -> None:
        received: dict[str, float] = {}
        for token in line.split()[1:]:
            if "=" not in token: continue
            key, value = token.split("=", 1)
            if key in self.vars:
                try:
                    received[key] = float(value)
                except ValueError: pass
        # During a table transaction the UI variables are the requested target.
        # Never overwrite them with a stale/partial device echo, otherwise a
        # retry would repeatedly resend the wrong value forever.
        if not self.table_waiting:
            for key, value in received.items():
                self.vars[key].set(value)
            self.draw_preview()
        if self.table_waiting:
            expected = dict(DEFAULTS)
            expected.update({k: float(v) for k, v in self.table_rows[self.table_index].items() if k in DEFAULTS})
            required = ("HMAX", "HC", "BR", "BS", "LOOP_HZ") if partial else tuple(expected)
            missing_required = [key for key in required if key not in received]
            different = [key for key, value in expected.items()
                         if key in received and abs(received[key] - value) > max(0.002, abs(value) * 0.0002)]
            if missing_required or different:
                details = []
                if missing_required: details.append("缺少 " + ",".join(missing_required))
                if different:
                    details.append("不一致 " + ", ".join(
                        f"{key}:期望{expected[key]:g}/收到{received[key]:g}" for key in different))
                self.append_log("STATUS 尚不能确认（" + "；".join(details) + "），继续等待")
                if self.status_retry_job:
                    self.after_cancel(self.status_retry_job)
                retry_delay_ms = min(5000, 400 * (2 ** min(max(self.table_apply_attempts - 1, 0), 4)))
                self.status_retry_job = self.after(
                    retry_delay_ms, lambda index=self.table_index: self._retry_apply_if_waiting(index))
                return
            if partial:
                missing = [key for key in expected if key not in received]
                self.append_log(
                    "旧固件 STATUS 被 100ms 发送超时截断；已核对收到的字段一致并继续"
                    + ("（未回传 " + ",".join(missing) + "）" if missing else "")
                )
            if self.status_retry_job:
                self.after_cancel(self.status_retry_job)
                self.status_retry_job = None
            if self.apply_watchdog_job:
                self.after_cancel(self.apply_watchdog_job)
                self.apply_watchdog_job = None
            self.table_waiting = False
            self.table_position.set(f"{self.table_index + 1} / {len(self.table_rows)}（设备已应用）")
            if self.capture_running:
                delay = max(0.0, float(self.capture_settle.get()))
                self.capture_status.set(f"参数已确认，等待示波器稳定 {delay:g} 秒…")
                self.table_timer = self.after(int(delay * 1000), self._capture_current_row)
            else:
                self.schedule_next_table_row()

    def _retry_apply_if_waiting(self, expected_index: int) -> None:
        self.status_retry_job = None
        if not self.table_waiting or not self.ser or self.table_index != expected_index:
            return
        self.table_apply_attempts += 1
        delay_note = ""
        if self.table_apply_attempts > 5:
            delay_note = "；蓝牙链路持续异常，仍会自动重试"
        self.capture_status.set(
            f"第 {expected_index + 1} 行参数不一致，正在重新下发 "
            f"(第 {self.table_apply_attempts} 次){delay_note}")
        self.append_log(
            f"重新下发第 {expected_index + 1} 行完整参数 "
            f"(第 {self.table_apply_attempts} 次)")
        self.discard_pending_tx()
        self.apply_expected_table_row()

    def append_log(self, text: str) -> None:
        self.log.config(state="normal")
        self.log.insert("end", text + "\n")
        self.log.see("end")
        self.log.config(state="disabled")

    def draw_preview(self) -> None:
        c = self.canvas; c.delete("all")
        w, h = max(c.winfo_width(), 200), max(c.winfo_height(), 200)
        pad = 35; c.create_line(pad, h/2, w-pad, h/2, fill="#52606d"); c.create_line(w/2, pad, w/2, h-pad, fill="#52606d")
        try:
            p = {k: v.get() for k, v in self.vars.items()}; ratio = min(max(p["BR"]/p["BS"], .001), .999)
            shape = p["HC"]/math.atanh(ratio); raw=[]
            for i in range(512):
                ph=2*math.pi*i/512; hv=p["HMAX"]*math.sin(ph)
                bv=p["BS"]*math.tanh((hv+p["HC"] if math.cos(ph)>=0 else hv-p["HC"])/shape)
                bv *= 1+p["ASYMMETRY"] if math.cos(ph)>=0 else 1-p["ASYMMETRY"]
                xn=p["X_GAIN"]*hv/p["HMAX"]+p["X_OFFSET"]
                yn=p["Y_GAIN"]*bv/p["BS"]+p["Y_OFFSET"]+p["XY_COUPLING"]*hv/p["HMAX"]
                raw.append((xn, yn))
            raw.append(raw[0])
            lengths=[0.0]
            for i in range(1, len(raw)):
                lengths.append(lengths[-1]+math.hypot(raw[i][0]-raw[i-1][0], raw[i][1]-raw[i-1][1]))
            pts=[]; segment=1
            for i in range(512):
                target=lengths[-1]*i/512
                while segment < 512 and lengths[segment] < target: segment += 1
                span=lengths[segment]-lengths[segment-1]
                mix=(target-lengths[segment-1])/span if span > 1e-12 else 0.0
                xn=raw[segment-1][0]+(raw[segment][0]-raw[segment-1][0])*mix
                yn=raw[segment-1][1]+(raw[segment][1]-raw[segment-1][1])*mix
                pts += [w/2+xn*(w/2-pad), h/2-yn*(h/2-pad)]
            c.create_line(*pts, fill="#38bdf8", width=2, smooth=False)
            c.create_text(12, 12, anchor="nw", fill="#d9e2ec", text=f"理论预览  {p['LOOP_HZ']:.1f} Hz")
        except Exception: pass

    def close(self) -> None:
        self.capture_stop.set(); self.stop_table_timer(); self.disconnect(); self.destroy()


if __name__ == "__main__":
    HysteresisUI().mainloop()
